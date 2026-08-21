# -*- coding: utf-8 -*-
"""Recalcula a pasta no LibreOffice e reporta erros de fórmula.

Usa a conversão direta (--convert-to), e não o despacho de macro Basic: em
arquivos grandes o caminho da macro fica preso esperando o documento carregar.
Como a pasta é gravada com fullCalcOnLoad, o próprio carregamento já recalcula
todas as fórmulas; a conversão apenas grava o resultado de volta.
"""
import os
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, "/root/.claude/skills/synced/xlsx/scripts")
from office.soffice import get_soffice_env                      # noqa: E402

from openpyxl import load_workbook                              # noqa: E402

ERROS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NULL!", "#NUM!",
         "Err:", "#ERRO", "#VALOR!", "#NOME?")


def recalcular(caminho, timeout=3600):
    caminho = os.path.abspath(caminho)
    saida = tempfile.mkdtemp(prefix="recalc_out_")
    perfil = tempfile.mkdtemp(prefix="recalc_prof_")
    env = get_soffice_env()
    cmd = ["soffice", "--headless", "--norestore",
           "-env:UserInstallation=file://" + perfil,
           "--convert-to", "xlsx:Calc MS Excel 2007 XML",
           "--outdir", saida, caminho]
    r = subprocess.run(cmd, env=env, capture_output=True, timeout=timeout)
    gerado = os.path.join(saida, os.path.basename(caminho))
    if not os.path.exists(gerado):
        raise SystemExit("LibreOffice não gerou saída. rc={} err={}".format(
            r.returncode, r.stderr.decode()[:400]))
    shutil.move(gerado, caminho)
    shutil.rmtree(saida, ignore_errors=True)
    shutil.rmtree(perfil, ignore_errors=True)
    return caminho


def conferir_erros(caminho):
    wb = load_workbook(caminho, data_only=True, read_only=True)
    achados, total = {}, 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=False):
            for c in row:
                v = c.value
                if isinstance(v, str) and any(v.startswith(e) for e in ERROS):
                    total += 1
                    k = (ws.title, v)
                    if k not in achados:
                        achados[k] = c.coordinate
    wb.close()
    return total, achados


if __name__ == "__main__":
    alvo = sys.argv[1]
    tl = int(sys.argv[2]) if len(sys.argv) > 2 else 3600
    import time
    t = time.time()
    recalcular(alvo, tl)
    n, ach = conferir_erros(alvo)
    print("recalculado em {:.0f}s · células com erro: {}".format(time.time() - t, n))
    for (aba, err), cel in list(ach.items())[:20]:
        print("   {} {} em {}".format(aba, err, cel))
    sys.exit(1 if n else 0)
