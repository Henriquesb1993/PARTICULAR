# -*- coding: utf-8 -*-
"""Confere a planilha gerada contra um recalculo independente feito em Python."""
import csv, datetime as dt, sys, collections
from openpyxl import load_workbook

VLR_TIT, VLR_AGR = 151.73, 544.95
COMBO = {0: 0.0, 1: 239.93, 2: 446.62, 3: 579.83, 4: 663.57, 5: 846.86}
LIM = 25


def d(s):
    for f in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s.strip(), f).date()
        except ValueError:
            pass
    return None


def idade(nasc, base):
    return base.year - nasc.year - ((base.month, base.day) < (nasc.month, nasc.day))


def norm(s):
    s = s.upper().strip()
    for a, b in [("Á","A"),("À","A"),("Â","A"),("Ã","A"),("Ä","A"),("É","E"),("Ê","E"),
                 ("Í","I"),("Ó","O"),("Ô","O"),("Õ","O"),("Ú","U"),("Ü","U"),("Ç","C"),
                 (".",""),("-"," ")]:
        s = s.replace(a, b)
    return " ".join(s.split())


def main(csv_in, xlsx, limite=None):
    linhas = []
    with open(csv_in, encoding="utf-8-sig", newline="") as fh:
        r = csv.reader(fh, delimiter=";")
        cab = None
        for row in r:
            if not any(c.strip() for c in row):
                continue
            v = [c.strip() for c in row]
            if cab is None:
                if v[0].upper() == "CONTRATO":
                    cab = v
                continue
            linhas.append(dict(zip(cab, v)))
    if limite:
        linhas = linhas[:limite]

    wv = load_workbook(xlsx, data_only=True)
    base = wv["PARÂMETROS"]["C6"].value
    base = base.date() if isinstance(base, dt.datetime) else base
    print("Data-base da planilha:", base)

    # ---- modelo independente ----
    for x in linhas:
        x["_dt"] = d(x["nascimento"])
        x["_id"] = idade(x["_dt"], base)
        p = x["parentesco"].upper()
        x["_tipo"] = "TITULAR" if p == "TITULAR" else ("AGREGADO" if p == "AGREGADO" else "DEPENDENTE")
        x["_eleg"] = (x["_tipo"] == "DEPENDENTE" and x["_id"] > LIM and not p.startswith("CONJ"))
        x["_dup"] = None

    vistos = set()
    for x in linhas:
        k = (x["matricula"], x["_dt"], norm(x["beneficiario"]))
        x["_dup"] = k in vistos
        vistos.add(k)

    g = collections.OrderedDict()
    for x in linhas:
        g.setdefault(x["matricula"], []).append(x)

    tit = {}
    for m, rs in g.items():
        t = [x for x in rs if x["_tipo"] == "TITULAR" and not x["_dup"]]
        if not t:
            continue
        deps = [x for x in rs if x["_tipo"] == "DEPENDENTE" and not x["_dup"]]
        agrs = [x for x in rs if x["_tipo"] == "AGREGADO" and not x["_dup"]]
        combo = COMBO[min(len(deps), 5)]
        tit[m] = dict(nome=t[0]["beneficiario"], contrato=t[0]["CONTRATO"],
                      ndep=len(deps), nagr=len(agrs), combo=combo,
                      vagr=len(agrs) * VLR_AGR,
                      total=VLR_TIT + combo + len(agrs) * VLR_AGR,
                      cobrado=sum(float(x["cobrado"].replace(".", "").replace(",", ".") or 0) for x in rs),
                      eleg=sum(1 for x in deps if x["_eleg"]))

    modelo = dict(tit)

    # ---- compara com a aba TITULARES ----
    ws = wv["TITULARES"]
    erros, n = [], 0
    for r in range(2, ws.max_row + 1):
        re_ = ws.cell(row=r, column=1).value
        if re_ in (None, ""):
            continue
        n += 1
        m = str(re_)
        if m not in tit:
            erros.append("RE {} aparece em TITULARES mas não no modelo".format(m))
            continue
        e = tit[m]
        chk = [("nome", ws.cell(row=r, column=2).value, e["nome"]),
               ("contrato", ws.cell(row=r, column=3).value, e["contrato"]),
               ("ndep", ws.cell(row=r, column=7).value, e["ndep"]),
               ("nagr", ws.cell(row=r, column=8).value, e["nagr"]),
               ("combo", round(ws.cell(row=r, column=17).value or 0, 2), round(e["combo"], 2)),
               ("vagr", round(ws.cell(row=r, column=18).value or 0, 2), round(e["vagr"], 2)),
               ("total", round(ws.cell(row=r, column=19).value or 0, 2), round(e["total"], 2)),
               ("cobrado", round(ws.cell(row=r, column=20).value or 0, 2), round(e["cobrado"], 2)),
               ("eleg", ws.cell(row=r, column=15).value, e["eleg"])]
        for campo, got, exp in chk:
            if got != exp:
                erros.append("RE {} · {}: planilha={!r} esperado={!r}".format(m, campo, got, exp))
        del tit[m]
    for m in tit:
        erros.append("RE {} está no modelo mas falta em TITULARES".format(m))

    print("Titulares conferidos: {}".format(n))
    print("Divergências: {}".format(len(erros)))
    for e in erros[:25]:
        print("  ✗", e)

    # ---- totais do resumo ----
    rg = wv["RESUMO GERENCIAL"]
    tot_mod = dict(
        titulares=n,
        dep=sum(1 for x in linhas if x["_tipo"] == "DEPENDENTE" and not x["_dup"]),
        agr=sum(1 for x in linhas if x["_tipo"] == "AGREGADO" and not x["_dup"]),
        dup=sum(1 for x in linhas if x["_dup"]),
    )
    print("\nModelo  → titulares={titulares} dependentes={dep} agregados={agr} duplicados={dup}".format(**tot_mod))
    print("Planilha→ titulares={} dependentes={} agregados={} duplicados={}".format(
        rg["B6"].value, rg["E6"].value, rg["H6"].value, rg["E20"].value))
    tot = sum(VLR_TIT + COMBO[min(t["ndep"], 5)] + t["nagr"] * VLR_AGR for t in modelo.values())
    print("Custo mensal total  planilha={!r}  modelo={:.2f}".format(rg["B9"].value, tot))
    print("Custo empresa       planilha={!r}  modelo={:.2f}".format(rg["E9"].value, n * VLR_TIT))
    if rg["B9"].value is not None and abs(rg["B9"].value - tot) > 0.01:
        erros.append("custo total diverge")
    if rg["E20"].value != tot_mod["dup"]:
        erros.append("duplicados: planilha={} modelo={}".format(rg["E20"].value, tot_mod["dup"]))
    if rg["B6"].value != n or rg["E6"].value != tot_mod["dep"] or rg["H6"].value != tot_mod["agr"]:
        erros.append("cartões do resumo divergem")
    print("\nTOTAL DE DIVERGÊNCIAS: {}".format(len(erros)))
    for e in erros[:25]:
        print("  ✗", e)
    return len(erros)


if __name__ == "__main__":
    lim = int(sys.argv[3]) if len(sys.argv) > 3 else None
    sys.exit(1 if main(sys.argv[1], sys.argv[2], lim) else 0)
